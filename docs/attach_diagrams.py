import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc_path = r'docs/Avance_de_proyecto_de_aula.docx'
doc = docx.Document(doc_path)

# Insert page break
p_break = doc.add_paragraph()
p_break.paragraph_format.page_break_before = True

title_p = doc.add_heading('Anexo: Diagramas de Arquitectura, Diseño y Modelado del Sistema', level=1)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

intro_p = doc.add_paragraph(
    'A continuación se presentan los diagramas formales y visuales del sistema desarrollados para el Marketplace Universitario UTB, '
    'especificando la arquitectura técnica, modelo entidad-relación de base de datos, casos de uso por roles, secuencias críticas de seguridad y ciclo de vida de los estados.'
)
intro_p.paragraph_format.space_after = Pt(12)

diagrams = [
    {
        'title': '1. Diagrama de Arquitectura y Componentes (C4 / Capas)',
        'img': r'docs/diagramas/renders/01_arquitectura_c4.png',
        'desc': 'Ilustra la separación de responsabilidades entre el cliente Flutter Web, el servidor FastAPI (Python 3.12), el sistema de persistencia (SQLAlchemy ORM + SQLite/PostgreSQL) y el almacenamiento local de archivos estáticos (/uploads).'
    },
    {
        'title': '2. Diagrama Entidad-Relación (DER / Modelo de Base de Datos)',
        'img': r'docs/diagramas/renders/02_modelo_entidad_relacion_der.png',
        'desc': 'Estructura relacional de las tablas del sistema: Usuarios, Publicaciones, Categorías, Compras, Mensajes y Reseñas, reflejando las restricciones de integridad y claves foráneas.'
    },
    {
        'title': '3. Diagrama de Casos de Uso (UML)',
        'img': r'docs/diagramas/renders/03_casos_de_uso_uml.png',
        'desc': 'Especifica las interacciones de los tres actores del sistema: Estudiante Comprador, Estudiante Vendedor y Moderador UTB.'
    },
    {
        'title': '4. Diagrama de Secuencia: Registro y Autenticación Segura (RF-01)',
        'img': r'docs/diagramas/renders/04_secuencia_autenticacion_jwt.png',
        'desc': 'Flujo de registro con validación obligatoria del dominio institucional @utb.edu.co, encriptación con Bcrypt y emisión de tokens JWT.'
    },
    {
        'title': '5. Diagrama de Secuencia: Compra Segura y Validación de Moderación (RF-05 & RF-08)',
        'img': r'docs/diagramas/renders/05_secuencia_compra_segura.png',
        'desc': 'Muestra la transacción atómica de compra: verificación del estado de la publicación, bloqueo de artículos suspendidos, transición a estado vendido y apertura automática del canal de mensajería para entrega.'
    },
    {
        'title': '6. Diagrama de Estados: Ciclo de Vida de una Publicación',
        'img': r'docs/diagramas/renders/06_estados_ciclo_vida_publicacion.png',
        'desc': 'Máquina de estados finitos que modela las transiciones de una publicación (Activa, Vendida, Pausada, Bloqueada y Eliminada).'
    }
]

for diag in diagrams:
    h = doc.add_heading(diag['title'], level=2)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    
    dp = doc.add_paragraph(diag['desc'])
    dp.paragraph_format.space_after = Pt(8)
    
    if os.path.exists(diag['img']):
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_after = Pt(14)
        run = img_p.add_run()
        run.add_picture(diag['img'], width=Inches(6.2))
    else:
        print('Image not found:', diag['img'])

out_path = r'docs/Avance_de_proyecto_de_aula_con_diagramas.docx'
doc.save(out_path)
print('Documento guardado exitosamente en:', out_path)
